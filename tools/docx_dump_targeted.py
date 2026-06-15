from __future__ import annotations

import sys
from pathlib import Path

from docx import Document


def main() -> int:
    for raw_path in sys.argv[1:]:
        path = Path(raw_path)
        doc = Document(path)
        print(f"=== {path} ===")
        print("-- paragraphs --")
        for index, paragraph in enumerate(doc.paragraphs):
            text = paragraph.text.strip()
            if text:
                print(f"[P{index:03}] ({paragraph.style.name}) {text}")
        print("-- tables --")
        for table_index, table in enumerate(doc.tables):
            print(f"[T{table_index:02}] rows={len(table.rows)} cols={len(table.columns)}")
            for row_index, row in enumerate(table.rows[:20]):
                cells = [" ".join(cell.text.split()) for cell in row.cells]
                print(f"  R{row_index:02}: {' | '.join(cells)}")
            if len(table.rows) > 20:
                print("  ...")
        print()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
