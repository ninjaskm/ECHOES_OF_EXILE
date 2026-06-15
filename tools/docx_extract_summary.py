from __future__ import annotations

import sys
from pathlib import Path

from docx import Document


KEYWORDS = (
    "MVP",
    "Rochatus",
    "dash",
    "pena",
    "feather",
    "HUD",
    "mana",
    "portal",
    "multiplayer",
    "Release Candidate",
    "Azurion",
    "Oricalum",
)


def main() -> int:
    for raw_path in sys.argv[1:]:
        path = Path(raw_path)
        doc = Document(path)
        print(f"=== {path} ===")
        print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)}")
        print("-- headings --")
        for paragraph in doc.paragraphs:
            style = paragraph.style.name if paragraph.style else ""
            text = paragraph.text.strip()
            if text and style.startswith("Heading"):
                print(f"{style}: {text}")
        print("-- keyword paragraphs --")
        for index, paragraph in enumerate(doc.paragraphs):
            text = paragraph.text.strip()
            if text and any(keyword.lower() in text.lower() for keyword in KEYWORDS):
                print(f"[{index}] {text[:500]}")
        print()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
