from __future__ import annotations

from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
DOCS = [
    ROOT / "ExiledWorld_GDD_Completo (6).docx",
    ROOT / "ExiledWorld_PlanoDeDesenvolvimento (1).docx",
]

REPLACEMENTS = {
    "Pendencia conhecida: multiplayer 2 jogadores ainda nao foi validado manualmente; ha cobertura automatizada parcial para targeting, scaling e recompensa de jogadores proximos.": "Pendência conhecida: multiplayer 2 jogadores ainda não foi validado manualmente; há cobertura automatizada parcial para targeting, scaling e recompensa de jogadores próximos.",
    "MVP em estado Release Candidate para solo: portal Tier 1, Portal Shard, Rochatus, dash por double-jump, HUD compacto, SaveSystem, drops e reset validados por teste automatizado e playtest solo.": "MVP em estado Release Candidate para solo: portal Tier 1, Portal Shard, Rochatus, dash por double-jump, HUD compacto, SaveSystem, drops e reset validados por teste automatizado e playtest solo.",
    "• Cooperativo de 2 a 4 jogadores; MVP solo validado, teste manual com 2 jogadores pendente": "• Cooperativo de 2 a 4 jogadores; MVP solo validado, teste manual com 2 jogadores pendente",
    "• Targeting MVP: boss troca alvo pelo jogador mais proximo, mas mantem alvo atual se ele estiver com 30% ou menos da vida maxima": "• Targeting MVP: boss troca alvo pelo jogador mais próximo, mas mantém alvo atual se ele estiver com 30% ou menos da vida máxima",
    "• Recompensa multiplayer: jogadores proximos ao boss recebem XP/mensagem; validacao manual 2p ainda pendente": "• Recompensa multiplayer: jogadores próximos ao boss recebem XP/mensagem; validação manual 2p ainda pendente",
    "Double-jump dash; avanco rapido na direcao do movimento": "Double-jump dash; avanço rápido na direção do movimento",
    "Critério de conclusão do MVP: 1 boss completo (Rochatus), 1 portal funcional, sistema de combate, atributos e HUD rodando sem crash. Status atual: Release Candidate solo validado; multiplayer local 2p pendente de playtest.": "Critério de conclusão do MVP: 1 boss completo (Rochatus), 1 portal funcional, sistema de combate, atributos e HUD rodando sem crash. Status atual: Release Candidate solo validado; multiplayer local 2p pendente de playtest.",
    "Status: Release Candidate para solo. `npm.cmd run check` passou com 110 testes, JS syntax OK e JSON OK.": "Status: Release Candidate para solo. npm.cmd run check passou com 110 testes, JS syntax OK e JSON OK.",
    "Implementado: portal Tier 1 com Portal Shard, Rochatus, ataques reutilizaveis, drops 6x Azurion + 20x Oricalum, dash double-jump custando 50 mana, HUD compacto, SaveSystem e reset MVP.": "Implementado: portal Tier 1 com Portal Shard, Rochatus, ataques reutilizáveis, drops 6x Azurion + 20x Oricalum, dash double-jump custando 50 mana, HUD compacto, SaveSystem e reset MVP.",
    "Pendente: validacao manual com 2 jogadores para troca de alvo, foco em alvo com <=30% de vida maxima e recompensa para jogadores proximos.": "Pendente: validação manual com 2 jogadores para troca de alvo, foco em alvo com <=30% de vida máxima e recompensa para jogadores próximos.",
    "HUD compacto mostra HP, MP e XP no actionbar/sidebar sem bugs visuais": "HUD compacto mostra HP, MP e XP no actionbar/sidebar sem bugs visuais",
    "Multiplayer: 2 jogadores ainda pendente de validacao manual; cobertura automatizada parcial para targeting e recompensas": "Multiplayer: 2 jogadores ainda pendente de validação manual; cobertura automatizada parcial para targeting e recompensas",
    "Criar wireframe da HUD principal (HP, MP, XP, acessorio ativo); MVP usa actionbar/sidebar compacto": "Criar wireframe da HUD principal (HP, MP, XP, acessório ativo); MVP usa actionbar/sidebar compacto",
    "Entregar o loop central: explorar -> portal -> boss -> loot -> evoluir. Status atual: MVP RC solo validado; proximo passo e Fase 2 / Tier 1 completo.": "Entregar o loop central: explorar -> portal -> boss -> loot -> evoluir. Status atual: MVP RC solo validado; próximo passo é Fase 2 / Tier 1 completo.",
    "FSM reutilizavel para IA de boss: IDLE -> COMBAT -> STAGGER -> DEAD, com estados auxiliares por boss": "FSM reutilizável para IA de boss: IDLE -> COMBAT -> STAGGER -> DEAD, com estados auxiliares por boss",
    "HUD compacto no actionbar/sidebar: HP, MP e XP com barras curtas": "HUD compacto no actionbar/sidebar: HP, MP e XP com barras curtas",
    "Rochatus — IA MVP completa com RollAttack, SpikeWaveAttack e QuakeAttack reutilizaveis": "Rochatus — IA MVP completa com RollAttack, SpikeWaveAttack e QuakeAttack reutilizáveis",
    "Multiplayer base: HP scaling e contratos automatizados; playtest manual 2p pendente": "Multiplayer base: HP scaling e contratos automatizados; playtest manual 2p pendente",
    "Magia Dash: consumo 50 mana, cooldown 2s, impulso na direcao do movimento": "Magia Dash: consumo 50 mana, cooldown 2s, impulso na direção do movimento",
    "Multiplayer sync: HP scaling e contratos automatizados; teste 2p manual pendente": "Multiplayer sync: HP scaling e contratos automatizados; teste 2p manual pendente",
    "Testes: `npm.cmd run check` verde; solo validado; multiplayer 2p pendente": "Testes: npm.cmd run check verde; solo validado; multiplayer 2p pendente",
    "Rochatus · 1 portal · Portal Shard · Azurion/Oricalum · dash double-jump 50 mana · HUD compacto · reset MVP · solo validado; multiplayer 2p pendente": "Rochatus · 1 portal · Portal Shard · Azurion/Oricalum · dash double-jump 50 mana · HUD compacto · reset MVP · solo validado; multiplayer 2p pendente",
    "Contratos automatizados para targeting/recompensas; playtest 2p como gate antes de declarar multiplayer validado.": "Contratos automatizados para targeting/recompensas; playtest 2p como gate antes de declarar multiplayer validado.",
}


def set_text(paragraph, text: str) -> None:
    for run in paragraph.runs:
        run.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = text
    else:
        paragraph.add_run(text)


def replace_in_paragraph(paragraph) -> None:
    text = paragraph.text
    if text in REPLACEMENTS:
        set_text(paragraph, REPLACEMENTS[text])


def replace_in_cell(cell) -> None:
    for paragraph in cell.paragraphs:
        replace_in_paragraph(paragraph)


def main() -> int:
    for path in DOCS:
        doc = Document(path)
        for paragraph in doc.paragraphs:
            replace_in_paragraph(paragraph)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    replace_in_cell(cell)

        if "PlanoDeDesenvolvimento" in path.name:
            if len(doc.tables) > 4:
                table = doc.tables[4]
                table.rows[1].cells[2].paragraphs[0].text = "src/scripts/core/EventBus.ts"
                table.rows[2].cells[2].paragraphs[0].text = "src/scripts/core/FSM.ts"
                table.rows[3].cells[2].paragraphs[0].text = "src/scripts/core/TickManager.ts"
                table.rows[4].cells[2].paragraphs[0].text = "src/scripts/save/SaveSystem.ts"
                table.rows[5].cells[2].paragraphs[0].text = "src/scripts/bosses/base/BaseBossSystem.ts"
            if len(doc.tables) > 5:
                table = doc.tables[5]
                if len(table.rows) == 4:
                    row = table.add_row()
                    row.cells[0].text = "exile:portal_structures"
                    row.cells[1].text = "JSON string"
                    row.cells[2].text = "[{ dimensionId, location }] — estruturas de portal geradas para reset MVP"
            if len(doc.tables) > 6:
                table = doc.tables[6]
                table.rows[7].cells[1].paragraphs[0].text = "Rochatus IA: FSM com estados IDLE, COMBAT, STAGGER, DEAD"

        if "GDD" in path.name:
            if len(doc.tables) > 16:
                doc.tables[16].rows[6].cells[1].paragraphs[0].text = "MVP: actionbar/sidebar compacto; futuro: Custom UI JSON + ActionForms"

        doc.save(path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
