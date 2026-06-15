from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
GDD = ROOT / "ExiledWorld_GDD_Completo (6).docx"
PLAN = ROOT / "ExiledWorld_PlanoDeDesenvolvimento (1).docx"


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    for run in paragraph.runs:
        run.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = text
    else:
        paragraph.add_run(text)


def insert_after(paragraph: Paragraph, text: str, style: str | None = None) -> Paragraph:
    new_paragraph = paragraph._parent.add_paragraph()
    if style:
        new_paragraph.style = style
    set_paragraph_text(new_paragraph, text)
    paragraph._p.addnext(new_paragraph._p)
    return new_paragraph


def set_cell_text(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    set_paragraph_text(paragraph, text)
    for extra in cell.paragraphs[1:]:
        set_paragraph_text(extra, "")


def update_gdd() -> None:
    doc = Document(GDD)

    # Current MVP snapshot near the title block.
    subtitle = next((p for p in doc.paragraphs if "Game Design Document" in p.text), None)
    if subtitle and not any("Estado atual do MVP" in p.text for p in doc.paragraphs):
        cursor = insert_after(subtitle, "Estado atual do MVP - v0.1.28", "normal")
        cursor = insert_after(
            cursor,
            "MVP em estado Release Candidate para solo: portal Tier 1, Portal Shard, Rochatus, dash por double-jump, HUD compacto, SaveSystem, drops e reset validados por teste automatizado e playtest solo.",
            "normal",
        )
        insert_after(
            cursor,
            "Pendencia conhecida: multiplayer 2 jogadores ainda nao foi validado manualmente; ha cobertura automatizada parcial para targeting, scaling e recompensa de jogadores proximos.",
            "normal",
        )

    replacements = {
        "▸ Rolamento — gira o corpo em alta velocidade na direção do jogador": "▸ Rolamento — gira em alta velocidade, trava a direcao inicial do alvo e empurra jogadores atingidos",
        "▸ Lançamento de espinhos — jogados pro céu, causam lentidão 3s ao cair (queda contínua 9s)": "▸ Queda de estalactites — particulas de estalactite caem em ondas, causam dano e lentidao ao impactar",
        "▸ Terremoto — pisa no chão liberando espinhos do solo": "▸ Terremoto — salta, impacta ao pousar, causa dano em area e camera shake",
        "• Cooperativo de 2 a 4 jogadores": "• Cooperativo de 2 a 4 jogadores; MVP solo validado, teste manual com 2 jogadores pendente",
        "• Histerese: debuff de buff de boss só acontece 30s após jogador sair — evita oscilação": "• Targeting MVP: boss troca alvo pelo jogador mais proximo, mas mantem alvo atual se ele estiver com 30% ou menos da vida maxima",
        "• Parry validado via System.run com prioridade — mitiga problemas de latência (ping ~100ms)": "• Recompensa multiplayer: jogadores proximos ao boss recebem XP/mensagem; validacao manual 2p ainda pendente",
    }
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text in replacements:
            set_paragraph_text(paragraph, replacements[text])

    # Tables: mana/dash, ores, tech, roadmap, risks.
    if len(doc.tables) > 8:
        set_cell_text(doc.tables[8].rows[1].cells[1], "50")
        set_cell_text(doc.tables[8].rows[1].cells[3], "Double-jump dash; avanco rapido na direcao do movimento")
    if len(doc.tables) > 10:
        set_cell_text(doc.tables[10].rows[1].cells[2], "Minerando; recompensa MVP atual inclui 6x Azurion ao derrotar Rochatus")
        set_cell_text(doc.tables[10].rows[2].cells[2], "Derrotar Rochatus; recompensa MVP atual inclui 20x Oricalum")
    if len(doc.tables) > 16:
        set_cell_text(doc.tables[16].rows[6].cells[1], "MVP: actionbar/sidebar compacto; futuro: Custom UI JSON + ActionForms")
    if len(doc.tables) > 19:
        set_cell_text(
            doc.tables[19].rows[2].cells[2],
            "Rochatus · 1 portal · Portal Shard · Azurion/Oricalum · dash double-jump 50 mana · HUD compacto · reset MVP · solo validado; multiplayer 2p pendente",
        )
    if len(doc.tables) > 20:
        set_cell_text(
            doc.tables[20].rows[3].cells[1],
            "Cobertura automatizada para targeting/recompensa; manter playtest 2p como gate antes de declarar multiplayer validado",
        )

    doc.save(GDD)


def update_plan() -> None:
    doc = Document(PLAN)

    replacements = {
        "Critério de conclusão do MVP: 1 boss completo (Rochatus), 1 portal funcional, sistema de combate, atributos e HUD rodando sem crash em multiplayer local.": "Critério de conclusão do MVP: 1 boss completo (Rochatus), 1 portal funcional, sistema de combate, atributos e HUD rodando sem crash. Status atual: Release Candidate solo validado; multiplayer local 2p pendente de playtest.",
        "HUD mostra HP, Mana e XP sem bugs visuais": "HUD compacto mostra HP, MP e XP no actionbar/sidebar sem bugs visuais",
        "Multiplayer: 2 jogadores testados localmente sem desync": "Multiplayer: 2 jogadores ainda pendente de validacao manual; cobertura automatizada parcial para targeting e recompensas",
        "Entregar o loop central: explorar → portal → boss → loot → evoluir. Um boss, um portal, sistemas funcionais. Estimativa: 3–4 semanas.": "Entregar o loop central: explorar -> portal -> boss -> loot -> evoluir. Status atual: MVP RC solo validado; proximo passo e Fase 2 / Tier 1 completo.",
        "Criar wireframe da HUD principal (HP, Mana, XP, acessório ativo)": "Criar wireframe da HUD principal (HP, MP, XP, acessorio ativo); MVP usa actionbar/sidebar compacto",
    }
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text in replacements:
            set_paragraph_text(paragraph, replacements[text])

    heading = next((p for p in doc.paragraphs if "MVP — Produto" in p.text), None)
    if heading and not any("Resumo atual do MVP" in p.text for p in doc.paragraphs):
        cursor = insert_after(heading, "Resumo atual do MVP - v0.1.28", "normal")
        cursor = insert_after(cursor, "Status: Release Candidate para solo. `npm.cmd run check` passou com 110 testes, JS syntax OK e JSON OK.", "normal")
        cursor = insert_after(cursor, "Implementado: portal Tier 1 com Portal Shard, Rochatus, ataques reutilizaveis, drops 6x Azurion + 20x Oricalum, dash double-jump custando 50 mana, HUD compacto, SaveSystem e reset MVP.", "normal")
        insert_after(cursor, "Pendente: validacao manual com 2 jogadores para troca de alvo, foco em alvo com <=30% de vida maxima e recompensa para jogadores proximos.", "normal")

    if len(doc.tables) > 2:
        table = doc.tables[2]
        set_cell_text(table.rows[2].cells[1], "FSM reutilizavel para IA de boss: IDLE -> COMBAT -> STAGGER -> DEAD, com estados auxiliares por boss")
        set_cell_text(table.rows[4].cells[1], "Mana base 200 · regen 5/s · dash double-jump custa 50 mana")
        set_cell_text(table.rows[7].cells[1], "HUD compacto no actionbar/sidebar: HP, MP e XP com barras curtas")
        set_cell_text(table.rows[8].cells[1], "Rochatus — IA MVP completa com RollAttack, SpikeWaveAttack e QuakeAttack reutilizaveis")
        set_cell_text(table.rows[10].cells[1], "Multiplayer base: HP scaling e contratos automatizados; playtest manual 2p pendente")
        set_cell_text(table.rows[11].cells[1], "Dash double-jump com custo de mana 50 e cooldown 2s; sem fallback por pena")
        set_cell_text(table.rows[13].cells[1], "Drop de Rochatus via loot table: 6x Azurion e 20x Oricalum")

    if len(doc.tables) > 6:
        table = doc.tables[6]
        status_done = "[x] Feito"
        status_partial = "[~] Parcial"
        for row_index in range(1, min(len(table.rows), 18)):
            set_cell_text(table.rows[row_index].cells[3], status_done)
        set_cell_text(table.rows[13].cells[1], "HUD compacto — actionbar/sidebar com HP, MP e XP")
        set_cell_text(table.rows[14].cells[1], "Magia Dash: consumo 50 mana, cooldown 2s, impulso na direcao do movimento")
        set_cell_text(table.rows[16].cells[1], "Multiplayer sync: HP scaling e contratos automatizados; teste 2p manual pendente")
        set_cell_text(table.rows[16].cells[3], status_partial)
        set_cell_text(table.rows[18].cells[1], "Testes: `npm.cmd run check` verde; solo validado; multiplayer 2p pendente")
        set_cell_text(table.rows[18].cells[3], status_partial)

    if len(doc.tables) > 7:
        table = doc.tables[7]
        for row_index in range(1, min(len(table.rows), 17)):
            set_cell_text(table.rows[row_index].cells[3], "[~] Placeholder/Parcial")
        set_cell_text(table.rows[10].cells[3], "[x] Feito")
        set_cell_text(table.rows[12].cells[1], "HUD — MVP compacto em actionbar/sidebar; UI JSON final fica para polish")
        set_cell_text(table.rows[12].cells[3], "[x] Feito MVP")
        set_cell_text(table.rows[14].cells[3], "[x] Feito")

    if len(doc.tables) > 12:
        set_cell_text(
            doc.tables[12].rows[2].cells[2],
            "Rochatus IA completa, Portal Tier 1, Combat, Mana, XP, HUD compacto, Save, reset MVP, solo validado; multiplayer 2p pendente",
        )
    if len(doc.tables) > 13:
        set_cell_text(
            doc.tables[13].rows[4].cells[2],
            "Contratos automatizados para targeting/recompensas; playtest 2p como gate antes de declarar multiplayer validado.",
        )

    doc.save(PLAN)


def main() -> int:
    update_gdd()
    update_plan()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
